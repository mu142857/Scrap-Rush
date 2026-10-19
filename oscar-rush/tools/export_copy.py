"""把全部文案导出成 Excel，给人手动改（改完用 tools/import_copy.py 导回）。

    python3 tools/export_copy.py              # 默认写到 ../文案/全部文案_给同学修改.xlsx
    python3 tools/export_copy.py 别的路径.xlsx

三页：
- 说明：给改文案的人看的规则 + 一行示例；
- 全部文案：每行一条 text.csv 的文字，按场景分组。只有黄色的"改后英文 / 改后中文"两列给人填；
- 写死在代码里：不走 text.csv 的给玩家看的字（人名、开机画面道具等），要改得改代码。
"英文原文"那列就是 text.csv 的 key，导回时靠它对行，别改。
"""
import collections
import csv
import re
import sys
from pathlib import Path

from fontTools.ttLib import TTFont
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from check_text import IGNORE, code_strings  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
CSV = ROOT / "localization/text.csv"
FONT = ROOT / "fonts/PixelFont.ttf"
DEFAULT_OUT = ROOT.parent / "文案/全部文案_给同学修改.xlsx"

## [分组名, 文件路径里含这些就归这一组]，按顺序第一个对上的算
GROUPS = [
    ("开始菜单", ["ui/title_screen", "ui/intro", "ui/mode_select", "ui/level_select", "ui/difficulty_select", "ui/confirm_box"]),
    ("设置和按键", ["ui/settings_menu", "ui/key_bind_menu", "autoload/controls", "autoload/settings"]),
    ("暂停", ["ui/pause_menu"]),
    ("教学关", ["tutorial/"]),
    ("序章、幕间、关卡开场", ["ui/prologue", "ui/interlude", "ui/level_intro", "levels/l3_boot_intro", "levels/l3_mode"]),
    ("关卡里的 HUD 和提示", ["ui/hud", "ui/coach", "ui/emotion_fx", "crystal/", "ui/hint_sign", "levels/encounter", "levels/level2",
                         "levels/level3.gd", "levels/test_level", "levels/tube_run", "autoload/run", "audio/", "ui/achievement_toast",
                         "combat/weapon_pickup"]),
    ("商店", ["shop/"]),
    ("武器", ["ui/weapon_select", "combat/weapon_registry", "ui/codex_weapons"]),
    ("变异和装备", ["ui/codex_mutations", "combat/mutations"]),
    ("怪物图鉴", ["ui/bestiary", "ui/codex_menu"]),
    ("人物和对白", ["ui/codex_crew", "ui/crew_line", "ui/dialog_box"]),
    ("成就", ["ui/achievements", "ui/codex_achievements"]),
    ("Boss 和精英", ["enemies/", "levels/level3_mobs", "boss/", "levels/boss_room", "levels/boss2_room", "ui/boss_bar"]),
    ("结局、结算和制作人员", ["levels/boss3_room", "ui/results_screen", "ui/credits_menu"]),
]
COMPOSED = "代码里拼出来的"
UNUSED = "可能已经不用了（可以跳过）"

## 这些文件里的字放在多宽的框里（11 号字时的像素宽，一个汉字 11 像素）。没写的按"别比现在长太多"算
WIDTH = {
    "ui/coach": 440, "ui/hud": 470, "shop/shop_room": 222, "shop/shop_stand": 80, "levels/encounter": 120,
    "tutorial/tutorial_level": 360, "ui/prologue": 360, "ui/interlude": 360, "ui/pause_menu": 470,
    "ui/title_screen": 470, "ui/settings_menu": 470, "levels/level2": 240, "levels/level3.gd": 240,
    "levels/test_level": 240, "ui/emotion_fx": 470, "ui/results_screen": 300,
}
## 没有固定框的，允许比现在最长的那种语言再长这么多
SLACK = 1.3

## 写死在代码里、不走 text.csv 的给玩家看的字：[文字, 在哪, 说明]
HARDCODED = [
    ("Mumu / Oscar / Little Oscar / Zpbrqols / Wutong / Ben2tu / Bentootoo / Rouri", "制作人员（credits_menu.gd）", "人名，两种语言都不翻"),
    ("SCRAP-TEK 86   BIOS v1.986", "第三关开机画面（l3_boot_intro.gd）", "老电脑开机画面的品牌行，画面道具"),
    ("%dK（内存计数）、OK", "第三关开机画面（l3_boot_intro.gd）", "老电脑自检的数字和 OK"),
    ("CH%d", "老电视 Boss 屏幕左上角的频道号（tv_boss.gd）", "自己画的 3×5 小字，只认 C、H 和数字"),
    ("WASD、Q/E、Tab、LB/RB、LT/RT、A/B/X/Y、LS/RS、Share", "各种按键提示", "按键名，两种语言一样，会自动替换进 %s"),
    ("开发者面板（F8）的全部文字", "scripts/debug/dev_mode.gd", "开发用，不给玩家看，只有英文"),
    ("启动图 / 浏览器加载页上的游戏名", "art/ui/boot_splash.png", "画在图片里，要改得重画"),
]


def group_of(files: list[str]) -> str:
    for name, pats in GROUPS:
        if any(p in f for f in files for p in pats):
            return name
    return "其他"


def width_of(files: list[str]):
    ws = [w for f in files for p, w in WIDTH.items() if p in f]
    return min(ws) if ws else None


class Measure:
    def __init__(self) -> None:
        f = TTFont(FONT)
        self.cmap = f.getBestCmap()
        self.hmtx = f["hmtx"]
        self.upm = f["head"].unitsPerEm

    def px(self, s: str) -> float:
        """11 号字时最宽那一行的像素宽"""
        best = 0.0
        for line in s.split("\n"):
            w = sum(self.hmtx[self.cmap[ord(c)]][0] for c in line if ord(c) in self.cmap) * 11 / self.upm
            best = max(best, w)
        return best


def notes_for(en: str, zh: str) -> str:
    out = []
    ph = re.findall(r"%[ds]|\{\w+\}", en)
    if ph:
        out.append("含占位符 " + " ".join(dict.fromkeys(ph)) + "，个数和顺序都别动（%s 会换成按键名或名字，%d 是数字）")
    if "\n" in en or "\n" in zh:
        out.append("含换行（单元格里的折行），行数尽量不变")
    if re.search(r"  ", en):
        out.append("中间的连续空格是排版用的，别删")
    if en == zh:
        out.append("两种语言一样（按键名、品牌、缩写）")
    return "；".join(out)


def main() -> None:
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_OUT
    with open(CSV, encoding="utf-8", newline="") as f:
        rows = list(csv.reader(f))[1:]
    keys = {r[0] for r in rows}
    loc: dict[str, list] = collections.defaultdict(list)
    for gd, line, s in code_strings():
        if s in keys:
            rel = gd.relative_to(ROOT).as_posix()
            loc[s].append((rel, line))
    m = Measure()

    items = []
    for i, (en, zh, *_) in enumerate(rows):
        where = loc.get(en, [])
        files = [w[0] for w in where]
        if where:
            grp = group_of(files)
        elif re.match(r"(TAPE \d|LEVEL \d|PAD\d)", en):
            grp = COMPOSED
        else:
            grp = UNUSED
        limit = width_of(files)
        cur = max(m.px(en), m.px(zh))
        if limit is None:
            cap = round(cur * SLACK)
            hint = f"别比现在长太多：最长约 {cap} 像素（≈{max(1, cap // 11)} 个汉字一行）"
        else:
            cap = limit
            hint = f"框宽约 {cap} 像素（≈{cap // 11} 个汉字 / {cap // 6} 个英文字母一行）"
        first = sorted(where)[0] if where else ("", 0)
        items.append({
            "grp": grp, "en": en, "zh": zh, "hint": hint, "cap": cap, "note": notes_for(en, zh),
            "where": "、".join(f"{f}:{ln}" for f, ln in sorted(set(where))[:4]),
            "order": ([g for g, _ in GROUPS] + [COMPOSED, UNUSED, "其他"]).index(grp), "first": first, "i": i,
        })
    items.sort(key=lambda d: (d["order"], d["first"], d["i"]))

    wb = Workbook()
    base = Font(name="Arial", size=11)
    bold = Font(name="Arial", size=11, bold=True)
    head_fill = PatternFill("solid", fgColor="2F2F2F")
    head_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    edit_fill = PatternFill("solid", fgColor="FFF2A8")
    thin = Side(style="thin", color="CCCCCC")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    # —— 说明 ——
    ws = wb.active
    ws.title = "说明"
    lines = [
        ("全部文案：给同学修改", bold),
        ("", base),
        ("这份表是游戏《Rewind Rush / 倒带冲刺》里给玩家看的全部文字，英文和中文各一份。", base),
        ("请在「全部文案」那一页修改，规则：", bold),
        ("1. 只填黄色的两列：「改后英文」「改后中文」。不想改的那一格留空就行，留空 = 不改。", base),
        ("2. 「英文原文」「中文」两列别动，导回时靠「英文原文」认是哪一句。", base),
        ("3. 占位符原样保留：%s、%d、{jump} 这种，个数和顺序都别变（%s 会被换成按键名或名字，%d 是数字）。", base),
        ("4. 单元格里的换行（Alt+Enter / Option+Enter 打出来的折行）是故意的，行数尽量别变；连续的几个空格也是排版用的，别删。", base),
        ("5. 别写太长：「最大长度提示」一列写了这句大约能放多宽，像素字体里一个汉字 11 像素、一个英文字母约 6 像素。太长会挤出框。", base),
        ("6. 不确定、想讨论的，写在「备注」里。", base),
        ("7. 语气：复古老科技、有点诗意，短句为主。精英怪的小标题可以俏皮（像老广告词），Boss 和结局安静一点。", base),
        ("8. 人名、按键名不翻；Rewind Rush 的中文名是「倒带冲刺」。", base),
        ("9. 分组「可能已经不用了」的那些可以跳过；「写死在代码里」那一页的字改不了表，有想法写下来就好。", base),
        ("", base),
        ("示例（假设想改一句提示）：", bold),
    ]
    for r, (text, font) in enumerate(lines, 1):
        c = ws.cell(row=r, column=1, value=text)
        c.font = font
    ex_head = ["英文原文", "中文", "改后英文", "改后中文", "备注"]
    ex_row = ["Go right", "往右走", "This way", "这边走", "（示例）想让路牌更口语一点"]
    r0 = len(lines) + 1
    for j, (h, v) in enumerate(zip(ex_head, ex_row), 1):
        hc = ws.cell(row=r0, column=j, value=h)
        hc.font, hc.fill, hc.border = head_font, head_fill, box
        vc = ws.cell(row=r0 + 1, column=j, value=v)
        vc.font, vc.border = base, box
        if h.startswith("改后"):
            vc.fill = edit_fill
    ws.column_dimensions["A"].width = 30
    for col in "BCDE":
        ws.column_dimensions[col].width = 22

    # —— 全部文案 ——
    ws = wb.create_sheet("全部文案")
    headers = ["编号", "出现在哪", "英文原文", "中文", "改后英文", "改后中文", "最大长度提示", "备注", "代码位置", "宽度上限(像素)"]
    widths = [7, 20, 48, 40, 40, 36, 30, 36, 34, 10]
    for j, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font, c.fill, c.border = head_font, head_fill, box
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = w
    for r, d in enumerate(items, 2):
        vals = [r - 1, d["grp"], d["en"], d["zh"], None, None, d["hint"], d["note"], d["where"], d["cap"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j)
            c.value = v
            if isinstance(v, str):
                c.data_type = "s"  # 以 = 开头的也当文字，别变成公式
            c.font, c.border, c.alignment = base, box, wrap
            if j in (5, 6):
                c.fill = edit_fill
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(items) + 1}"
    ws.column_dimensions["J"].hidden = True

    # —— 写死在代码里 ——
    ws = wb.create_sheet("写死在代码里")
    for j, (h, w) in enumerate(zip(["文字", "在哪", "说明", "想改成（可选）"], [52, 40, 46, 36]), 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font, c.fill, c.border = head_font, head_fill, box
        ws.column_dimensions[get_column_letter(j)].width = w
    for r, (t, w, n) in enumerate(HARDCODED, 2):
        for j, v in enumerate([t, w, n, None], 1):
            c = ws.cell(row=r, column=j, value=v)
            c.font, c.border, c.alignment = base, box, wrap
            if j == 4:
                c.fill = edit_fill

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    cnt = collections.Counter(d["grp"] for d in items)
    print(f"{out}：{len(items)} 条")
    for g, n in sorted(cnt.items(), key=lambda kv: ([x for x, _ in GROUPS] + [COMPOSED, UNUSED, "其他"]).index(kv[0])):
        print(f"  {g}: {n}")


if __name__ == "__main__":
    main()
