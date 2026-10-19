"""把同学改完的文案表（tools/export_copy.py 导出的 xlsx）写回游戏。

    python3 tools/import_copy.py 改好的表.xlsx            # 只检查，列出要改什么，不写任何文件
    python3 tools/import_copy.py 改好的表.xlsx --apply    # 检查没有错误才写回

需要 openpyxl 和 fontTools（系统 Python 不让装包时，建一个虚拟环境：
python3 -m venv /tmp/venv && /tmp/venv/bin/pip install openpyxl fonttools，再用 /tmp/venv/bin/python 跑）。

怎么认行：靠"英文原文"列（= text.csv 的 key），"编号"和顺序随便。
- 只填了"改后中文"：改 text.csv 第二列。
- 填了"改后英文"：英文原文就是代码里的字符串，要两边一起改——把 scripts/ 里所有一模一样的字符串字面量
  （双引号、单引号都算）换成新英文，再改 text.csv 的 key。代码里找不到原文（比如是拼出来的）就不改英文，只报告。
检查（有"错误"就不写回）：
- 错误：英文原文在 text.csv 里找不到（原文列被改了）；改后英文和别的行撞了；占位符（%s %d {jump}）个数或顺序变了；
  用了完整字体里也没有的字（游戏里会是方框）。
- 提醒：比"宽度上限"宽（可能挤出框）；换行数变了；改了第三关怪物名（level3_mobs.gd 是生成的，生成脚本也要改）。
写回以后自动重跑 tools/subset_font.py 和 tools/check_text.py。Godot 编辑器切回来会重新导入 text.csv；
命令行导出前跑一次 Godot --headless --import。
"""
import csv
import re
import subprocess
import sys
from pathlib import Path

from fontTools.ttLib import TTFont
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
CSV = ROOT / "localization/text.csv"
FULL_FONT = ROOT / "tools/fonts/PixelFont_full.ttf"
SHEET = "全部文案"
PLACEHOLDER = re.compile(r"%[ds]|\{\w+\}")


def norm(v) -> str:
    """单元格的值：None 当空；Excel 的换行统一成 \\n"""
    if v is None:
        return ""
    return str(v).replace("\r\n", "\n").replace("\r", "\n")


def literals(s: str) -> list[str]:
    """s 在 GDScript 里写成字面量的两种样子"""
    esc = s.replace("\\", "\\\\").replace("\n", "\\n").replace("\t", "\\t")
    return ['"' + esc.replace('"', '\\"') + '"', "'" + esc.replace("'", "\\'") + "'"]


class Width:
    def __init__(self) -> None:
        f = TTFont(FULL_FONT)
        self.cmap = f.getBestCmap()
        self.hmtx = f["hmtx"]
        self.upm = f["head"].unitsPerEm

    def px(self, s: str) -> float:
        best = 0.0
        for line in s.split("\n"):
            w = 0.0
            for c in line:
                if c == " ":
                    w += 3.0  # 子集里空格加宽到 3 像素（subset_font.py 的 SPACE_PX）
                elif ord(c) in self.cmap:
                    w += self.hmtx[self.cmap[ord(c)]][0] * 11 / self.upm
            best = max(best, w)
        return best

    def missing(self, s: str) -> str:
        return "".join(sorted({c for c in s if c not in "\n\t" and ord(c) not in self.cmap}))


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    xlsx = Path(sys.argv[1])
    apply = "--apply" in sys.argv[2:]

    with open(CSV, encoding="utf-8", newline="") as f:
        rows = list(csv.reader(f))
    header, body = rows[0], rows[1:]
    index = {r[0]: i for i, r in enumerate(body)}

    ws = load_workbook(xlsx)[SHEET]
    head = [norm(c.value) for c in ws[1]]
    col = {name: head.index(name) for name in ["英文原文", "改后英文", "改后中文"]}
    cap_col = head.index("宽度上限(像素)") if "宽度上限(像素)" in head else None

    gd_files = [p for p in (ROOT / "scripts").rglob("*.gd") if "debug" not in p.parts]
    gd_text = {p: p.read_text(encoding="utf-8") for p in gd_files}
    w = Width()

    errors, warns, zh_changes, en_changes = [], [], [], []
    new_keys: dict[str, str] = {}
    for r, cells in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        key = norm(cells[col["英文原文"]])
        new_en = norm(cells[col["改后英文"]])
        new_zh = norm(cells[col["改后中文"]])
        if not new_en.strip() and not new_zh.strip():
            continue
        tag = f"第 {r} 行 {key[:40]!r}"
        if key not in index:
            errors.append(f"{tag}：英文原文在 text.csv 里找不到（这一列别改）")
            continue
        old_zh = body[index[key]][1]
        en = new_en if new_en.strip() else key
        zh = new_zh if new_zh.strip() else old_zh
        want = PLACEHOLDER.findall(key)
        for label, text in (("改后英文", en), ("改后中文", zh)):
            if PLACEHOLDER.findall(text) != want:
                errors.append(f"{tag}：{label}的占位符 {PLACEHOLDER.findall(text)} 和原文 {want} 对不上（个数和顺序都要一样）")
            miss = w.missing(text)
            if miss:
                errors.append(f"{tag}：{label}里有字体没有的字 {miss!r}（游戏里会是方框，换个说法）")
            if text.count("\n") != key.count("\n"):
                warns.append(f"{tag}：{label}的行数和原文不一样（{key.count(chr(10)) + 1} → {text.count(chr(10)) + 1}）")
        cap = cells[cap_col] if cap_col is not None else None
        if isinstance(cap, (int, float)):
            for label, text in (("改后英文", en), ("改后中文", zh)):
                if w.px(text) > cap:
                    warns.append(f"{tag}：{label}约 {w.px(text):.0f} 像素，超过宽度上限 {cap}，可能挤出框")
        if new_zh.strip() and new_zh != old_zh:
            zh_changes.append((key, new_zh))
        if new_en.strip() and new_en != key:
            if new_en in index or new_en in new_keys:
                errors.append(f"{tag}：改后英文 {new_en!r} 和别的行一样了（英文是认行用的，不能重复）")
                continue
            hits = [(p, lit) for p in gd_files for lit in literals(key) if lit in gd_text[p]]
            if not hits:
                warns.append(f"{tag}：代码里找不到这句原文（可能是拼出来的），英文没改，只改了中文")
                continue
            if any(p.name == "level3_mobs.gd" for p, _ in hits):
                warns.append(f"{tag}：在 level3_mobs.gd 里，这个文件是 草图/第三关/怪物/build_mobs.py 生成的，生成脚本也要一起改")
            new_keys[new_en] = key
            en_changes.append((key, new_en, hits))

    print(f"读了 {xlsx.name}：改中文 {len(zh_changes)} 条，改英文 {len(en_changes)} 条")
    for key, zh in zh_changes:
        print(f"  中文  {key[:40]!r}：{body[index[key]][1]!r} → {zh!r}")
    for key, en, hits in en_changes:
        files = sorted({p.relative_to(ROOT).as_posix() for p, _ in hits})
        print(f"  英文  {key!r} → {en!r}（{', '.join(files)}）")
    for m in warns:
        print("提醒：" + m)
    for m in errors:
        print("错误：" + m)
    if errors:
        print(f"有 {len(errors)} 个错误，没有写回。改好表再跑。")
        return 1
    if not apply:
        print("只检查，没写文件。确认无误加 --apply 写回。")
        return 0

    for key, zh in zh_changes:
        body[index[key]][1] = zh
    for key, en, hits in en_changes:
        body[index[key]][0] = en
        for p, lit in hits:
            new_lit = literals(en)[0 if lit.startswith('"') else 1]
            gd_text[p] = gd_text[p].replace(lit, new_lit)
    for p in {p for _, _, hits in en_changes for p, _ in hits}:
        p.write_text(gd_text[p], encoding="utf-8")
    with open(CSV, "w", encoding="utf-8", newline="") as f:
        csv.writer(f, lineterminator="\n").writerows([header] + body)
    print("已写回 text.csv" + ("和代码" if en_changes else "") + "，重跑字体子集和检查：", flush=True)
    subprocess.run([sys.executable, str(ROOT / "tools/subset_font.py")], check=False)
    subprocess.run([sys.executable, str(ROOT / "tools/check_text.py")], check=False)
    if en_changes:
        print("改了代码：记得限时编译检查（Godot --headless --check-only 或导出一次），再在游戏里看一眼。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
